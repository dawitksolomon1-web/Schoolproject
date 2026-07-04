"""Best-effort plain-text extraction from downloaded course material files."""
from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger("canvas_assistant.text_extraction")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix == ".pptx":
            return _extract_pptx(path)
        if suffix in {".txt", ".md", ".csv"}:
            return path.read_text(errors="ignore")
    except Exception:
        logger.exception("Failed to extract text from %s", path)
        return ""
    logger.debug("No text extractor for %s, skipping", path)
    return ""


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    chunks = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks)
